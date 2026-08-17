#!/usr/bin/env python3
"""Build the resumable planner-pet image generation manifest.

The canonical workbook owns pet identity and count. The manifest owns only art
production state; it never writes gameplay data back to the workbook.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = REPO_ROOT.parent / "ysbzs" / "xlsx" / "ysbzs_master.xlsx"
DEFAULT_MANIFEST = (
    REPO_ROOT / "art" / "manifests" / "shared" / "pets" / "pet_image_generation_manifest.json"
)
OUTPUT_DIR = REPO_ROOT / "art" / "images" / "shared" / "pets" / "generated"
PET_FIELDS = (
    "pet_id",
    "name",
    "element",
    "tier",
    "role",
    "body_size",
    "source_object_id",
    "mechanism_id",
    "shape_id",
)
STYLE_REFERENCES = (
    "res://art/images/route/three_choice_psd/route_portrait_shop.png",
    "res://art/images/route/three_choice_psd/route_portrait_event.png",
    "res://art/images/route/three_choice_psd/route_portrait_reward.png",
    "res://art/images/battle/runtime/hero_images/hero_wukong.png",
    "res://art/images/battle/map_controls/maps/grassland_morning.png",
    "res://art/images/shared/pets/sheets/slices/pal_001.png",
)
STYLE_CONTRACT = {
    "id": "ysbzs-chibi-pixel-v1",
    "required": [
        "q-version creature proportions with a compact readable silhouette",
        "visible intentional pixel clusters and hard non-antialiased contour steps",
        "dark warm outer outline and a limited stepped-shading palette",
        "one complete centered creature on transparency without scenery or UI",
    ],
    "forbidden": [
        "photorealistic anatomy, fur, feathers, scales, or materials",
        "smooth painterly, vector, 3D-rendered, or airbrushed surfaces",
        "cinematic glow, depth-of-field blur, floor shadow, or background gradient",
        "card frame, ornamental portrait ring, text, logo, or watermark",
    ],
    "approval_rule": "identity review and style review must both pass before runtime mapping",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--manifest", type=Path, default=DEFAULT_MANIFEST)
    parser.add_argument("--check", action="store_true", help="Fail if the manifest is stale; do not write")
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def row_hash(record: dict[str, str]) -> str:
    payload = json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def load_pet_rows(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "PETS" not in workbook.sheetnames:
        raise ValueError(f"PETS sheet missing from {workbook_path}")
    rows = workbook["PETS"].iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    indexes = {name: headers.index(name) for name in PET_FIELDS}
    records: list[dict[str, str]] = []
    seen_ids: set[str] = set()
    for row in rows:
        pet_id = str(row[indexes["pet_id"]] or "").strip()
        if not pet_id:
            continue
        if pet_id in seen_ids:
            raise ValueError(f"duplicate pet_id in PETS: {pet_id}")
        seen_ids.add(pet_id)
        records.append({field: str(row[indexes[field]] or "").strip() for field in PET_FIELDS})
    if not records:
        raise ValueError("PETS has no data rows")
    return records


def prompt_for(record: dict[str, str]) -> str:
    return "\n".join(
        (
            "Use case: stylized-concept",
            "Asset type: reusable Godot pet sprite for route choice, shop, bag, party, and battle",
            (
                f"Primary request: create the unique pet {record['name']} ({record['pet_id']}), "
                f"a {record['body_size']} {record['role']} creature with {record['element']} affinity"
            ),
            "Input images: project route portraits, pixel battle hero, pixel battle map, and early pet sprite are binding style references; the pet slice is framing reference only",
            "Scene/backdrop: perfectly flat solid #ff00ff chroma-key background for removal",
            "Style/medium: authentic Chinese mythic Q-version game-sprite pixel art matching the references; clearly visible intentional pixel clusters, hard stair-step edges with no smooth brush blending, dark warm outline, limited stepped shading, compact readable silhouette",
            "Composition/framing: one complete creature, centered, three-quarter game-sprite view, full body visible, generous padding, no cropping",
            f"Lighting/mood: lively and approachable; visual importance consistent with {record['tier']} quality without adding a frame",
            "Constraints: identity must clearly match the Chinese pet name and affinity; one creature only; reusable without UI; no text, no logo, no watermark",
            "Avoid: card frame, ornamental ring, panel, scenery, floor plane, cast shadow, reflection, smooth vector or painterly art, 3D rendering, photorealistic anatomy/fur/feathers/scales, airbrushed gradients, cinematic glow, blur, or background texture",
            "Background removal constraints: uniform #ff00ff only; do not use #ff00ff on the creature; crisp separated silhouette",
        )
    )


def image_metadata(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {"state": "missing"}
    with Image.open(path) as image:
        bands = image.getbands()
        has_alpha = "A" in bands
        alpha_bbox = image.getchannel("A").getbbox() if has_alpha else None
        return {
            "state": "generated_unreviewed",
            "sha256": sha256_file(path),
            "width": image.width,
            "height": image.height,
            "mode": image.mode,
            "has_alpha": has_alpha,
            "alpha_bbox": list(alpha_bbox) if alpha_bbox else None,
        }


def load_previous(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def build_manifest(workbook_path: Path, manifest_path: Path) -> dict[str, Any]:
    records = load_pet_rows(workbook_path)
    previous_manifest = load_previous(manifest_path)
    previous = {
        str(item["pet_id"]): item for item in previous_manifest.get("targets", [])
    }
    style_audit = dict(previous_manifest.get("style_audit", {}))
    raw_needs_style_regeneration = [
        str(pet_id) for pet_id in style_audit.get("needs_regeneration_ids", [])
    ]
    needs_style_regeneration = set(raw_needs_style_regeneration)
    if len(raw_needs_style_regeneration) != len(needs_style_regeneration):
        raise ValueError("style audit contains duplicate needs_regeneration_ids")
    record_ids = {record["pet_id"] for record in records}
    unknown_style_ids = sorted(needs_style_regeneration - record_ids)
    if unknown_style_ids:
        raise ValueError(f"style audit contains unknown pet IDs: {unknown_style_ids}")
    audit_contract = str(style_audit.get("contract_version", ""))
    if audit_contract and audit_contract != STYLE_CONTRACT["id"]:
        raise ValueError(
            f"style audit contract {audit_contract!r} does not match {STYLE_CONTRACT['id']!r}"
        )
    targets: list[dict[str, Any]] = []
    for record in records:
        pet_id = record["pet_id"]
        expected_row_hash = row_hash(record)
        old = previous.get(pet_id, {})
        final_path = OUTPUT_DIR / f"{pet_id}.png"
        review_status = str(old.get("review_status", "pending"))
        review_note = str(old.get("review_note", ""))
        if old and old.get("planner_row_sha256") != expected_row_hash:
            review_status = "needs_regeneration"
            review_note = "planner identity fields changed after the previous art review"
        targets.append(
            {
                **record,
                "planner_row_sha256": expected_row_hash,
                "output_path": f"res://art/images/shared/pets/generated/{pet_id}.png",
                "prompt": prompt_for(record),
                "file": image_metadata(final_path),
                "review_status": review_status,
                "review_note": review_note,
            }
        )
    generated = sum(item["file"]["state"] != "missing" for item in targets)
    approved = sum(
        item["review_status"] == "approved"
        and item["pet_id"] not in needs_style_regeneration
        for item in targets
    )
    return {
        "schema": "ysbzs.pet-image-generation-manifest.v2",
        "source_workbook": "../ysbzs/xlsx/ysbzs_master.xlsx",
        "source_sheet": "PETS",
        "source_workbook_sha256": sha256_file(workbook_path),
        "style_references": list(STYLE_REFERENCES),
        "style_contract": STYLE_CONTRACT,
        "style_audit": style_audit,
        "target_count": len(targets),
        "generated_count": generated,
        "approved_count": approved,
        "pending_count": len(targets) - approved,
        "targets": targets,
    }


def canonical_json(value: dict[str, Any]) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2, sort_keys=False) + "\n"


def main() -> int:
    args = parse_args()
    manifest = build_manifest(args.workbook.resolve(), args.manifest.resolve())
    rendered = canonical_json(manifest)
    if args.check:
        current = args.manifest.read_text(encoding="utf-8") if args.manifest.is_file() else ""
        if current != rendered:
            print(f"PET_IMAGE_MANIFEST_STALE path={args.manifest}")
            return 1
        print(
            "PET_IMAGE_MANIFEST_OK "
            f"targets={manifest['target_count']} generated={manifest['generated_count']} "
            f"approved={manifest['approved_count']}"
        )
        return 0
    args.manifest.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    args.manifest.write_text(rendered, encoding="utf-8")
    print(
        "PET_IMAGE_MANIFEST_WRITTEN "
        f"targets={manifest['target_count']} generated={manifest['generated_count']} "
        f"approved={manifest['approved_count']} path={args.manifest}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
